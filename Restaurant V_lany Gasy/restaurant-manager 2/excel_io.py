"""Import / Export Excel.

Principes (pour que ça marche « à tous les coups ») :
- Les colonnes sont reconnues par leur NOM (en-tête), pas par leur position :
  l'ordre peut changer, les accents/majuscules/espaces ne comptent pas
  (« Nom ingrédient », « nom_ingredient », « NOM INGREDIENT » = pareil).
- Formats acceptés : .xlsx, .xlsm (Excel, LibreOffice, Google Sheets) et .csv.
- Nombres acceptés : 2500 · 2 500 · 0,3 · 0.3 · « 3000 Ar ».
- Dates acceptées : vraie date Excel, 23/09/2026, 23-09-2026, 2026-09-23.
- Un ingrédient peut être désigné par son NOM ou par son CODE (PLT, YRT...).
- Toutes les lignes sont vérifiées AVANT d'enregistrer quoi que ce soit : s'il y
  a une seule erreur, rien n'est importé et la liste des erreurs est affichée.
- Un fichier exporté peut être ré-importé tel quel sans rien dupliquer
  (ingrédients et menus). Le stock ré-importé REMPLACE le stock (et c'est tracé
  dans le journal), il ne s'ajoute pas.
"""
import csv
import io
import re
import unicodedata
from datetime import datetime

from flask import Blueprint, render_template, request, send_file, flash, redirect, url_for
from flask_login import login_required, current_user
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from extensions import db
from models import (
    Ingredient, MenuItem, RecipeItem, Purchase, FundingSource, Expense, Sale, StockAdjustment,
    EXPENSE_TYPE_FIXED, EXPENSE_TYPE_VARIABLE, ROLE_ADMIN,
)
from decorators import roles_required
from utils import (
    parse_float, clean_text, parse_date_any, now_local, to_local, local_to_utc,
    suggest_code, unique_code, normalize_code,
)

excel_bp = Blueprint("excel", __name__, url_prefix="/excel")

HEADER_FILL = PatternFill("solid", fgColor="6E3614")
HEADER_FONT = Font(bold=True, color="FFFFFF")
EXAMPLE_FONT = Font(italic=True, color="7A6A5E")


# ===========================================================================
# Outils communs
# ===========================================================================
def _norm(h):
    """« Quantité par portion » → « quantiteparportion »."""
    s = unicodedata.normalize("NFD", clean_text(h)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]", "", s)


# Nom logique de colonne → noms d'en-tête acceptés (déjà normalisés)
COLS = {
    "ingredients": {
        "name": ["nomingredient", "nom", "ingredient", "name", "designation", "libelle"],
        "code": ["code", "motcle", "codecourt", "abreviation"],
        "category": ["categorie", "category", "famille"],
        "unit": ["unite", "unit", "unites"],
        "stock": ["stock", "stockinitial", "stockactuel", "quantite", "quantitestock", "qte"],
        "threshold": ["seuilalerte", "seuil", "alerte", "stockmin", "stockminimum"],
        "price": ["prixunitaire", "prix", "pu", "prixunit", "prixachat"],
    },
    "menu": {
        "name": ["nommenu", "menu", "nom", "plat", "nomplat"],
        "category": ["categoriemenu", "categorie", "category"],
        "price": ["prixvente", "prix", "pv", "prixdevente"],
        "ingredient": ["ingredient", "codeingredient", "nomingredient", "ingredientoucode"],
        "qty": ["quantiteparportion", "quantite", "qte", "qteparportion", "dose"],
    },
    "achats": {
        "date": ["date", "dateachat"],
        "supplier": ["fournisseur", "supplier", "vendeur", "marche"],
        "ingredient": ["ingredient", "codeingredient", "nomingredient", "ingredientoucode"],
        "qty": ["quantite", "qte", "quantity"],
        "price": ["prixunitaire", "prix", "pu"],
        "source": ["sourcefinancement", "source", "financement", "caisse"],
    },
    "depenses": {
        "date": ["date"],
        "type": ["type", "typedepense", "fixeouvariable"],
        "category": ["categorie", "category", "nature"],
        "amount": ["montant", "amount", "somme", "total"],
        "description": ["description", "detail", "details", "commentaire", "note"],
        "source": ["sourcefinancement", "source", "financement", "caisse"],
    },
}
# En-têtes qui montrent qu'on s'est trompé de bouton (ex : fichier d'achats
# importé comme ingrédients → le stock serait écrasé par les quantités achetées)
FORBIDDEN = {
    "ingredients": {"nommenu", "prixvente", "fournisseur", "sourcefinancement", "montant", "quantiteparportion"},
    "menu": {"fournisseur", "montant", "stock", "seuilalerte", "sourcefinancement"},
    "achats": {"nommenu", "prixvente", "seuilalerte", "montant", "quantiteparportion"},
    "depenses": {"ingredient", "nommenu", "fournisseur", "seuilalerte", "quantite"},
}
KIND_LABELS = {"ingredients": "Ingrédients", "menu": "Menu + Recette", "achats": "Achats", "depenses": "Dépenses"}

REQUIRED = {
    "ingredients": ["name"],
    "menu": ["name"],
    "achats": ["ingredient", "qty"],
    "depenses": ["amount"],
}


class ImportFileError(Exception):
    pass


def _read_table(file_storage, kind):
    """Lit le fichier envoyé et retourne (mapping, lignes).
    mapping = {nom_logique: index_colonne} ; lignes = [(n°_ligne, [valeurs])].
    Cherche automatiquement la bonne feuille et la ligne d'en-têtes."""
    filename = (file_storage.filename or "").lower()
    raw = file_storage.read()
    if not raw:
        raise ImportFileError("Le fichier est vide.")

    sheets = []  # [(nom_feuille, [[valeurs]])]
    if filename.endswith(".csv") or filename.endswith(".txt"):
        text = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        sample = text[:2000]
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
        if sample.count("\t") > max(sample.count(";"), sample.count(",")):
            delimiter = "\t"
        sheets.append(("CSV", [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]))
    elif filename.endswith(".xls"):
        raise ImportFileError(
            "Ancien format .xls non supporté. Dans Excel : Fichier → Enregistrer sous → "
            "« Classeur Excel (.xlsx) », puis importez ce nouveau fichier."
        )
    else:
        try:
            wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        except Exception:
            raise ImportFileError(
                "Impossible de lire ce fichier. Vérifiez qu'il s'agit bien d'un fichier Excel "
                "(.xlsx) ou CSV, et qu'il n'est pas protégé par un mot de passe."
            )
        # Feuille active en premier, puis les autres
        ordered = [wb.active] + [ws for ws in wb.worksheets if ws is not wb.active]
        for ws in ordered:
            sheets.append((ws.title, [list(r) for r in ws.iter_rows(values_only=True)]))

    aliases = COLS[kind]
    for _title, rows in sheets:
        for h_idx, row in enumerate(rows[:15]):
            normed = [_norm(c) for c in row]
            mapping = {}
            for logical, names in aliases.items():
                for i, n in enumerate(normed):
                    if n and n in names and i not in mapping.values():
                        mapping[logical] = i
                        break
            if all(r in mapping for r in REQUIRED[kind]):
                wrong = FORBIDDEN[kind] & set(normed)
                if wrong:
                    guess = next((KIND_LABELS[k] for k, al in COLS.items() if k != kind
                                  and all(any(n in al[r] for n in normed) for r in REQUIRED[k])), None)
                    raise ImportFileError(
                        "Ce fichier ne correspond pas à cet import"
                        + (f" : il ressemble à un fichier « {guess} »" if guess else "")
                        + ". Utilisez le bon bouton « Importer » (rien n'a été modifié)."
                    )
                data = []
                for offset, r in enumerate(rows[h_idx + 1:], start=h_idx + 2):
                    if r is None or all(clean_text(c) == "" for c in r):
                        continue
                    first = clean_text(r[0]) if r else ""
                    if first.startswith("#") or first.lower().startswith("exemple"):
                        continue  # lignes de commentaire / d'exemple
                    data.append((offset, list(r)))
                return mapping, data

    expected = ", ".join(COLS[kind][c][0] for c in REQUIRED[kind])
    raise ImportFileError(
        f"En-têtes introuvables. La première ligne doit contenir au moins les colonnes : {expected}. "
        "Le plus simple : téléchargez le modèle, remplissez-le, puis importez-le."
    )


def _get(row, mapping, key):
    idx = mapping.get(key)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _ingredient_lookup():
    ings = Ingredient.query.all()
    by_name = {i.name.strip().lower(): i for i in ings}
    by_code = {i.code.upper(): i for i in ings if i.code}
    return by_name, by_code


def _find_ingredient(value, by_name, by_code):
    """Trouve un ingrédient par son nom exact (sans tenir compte des majuscules),
    sinon par son code (PLT, YRT...)."""
    txt = clean_text(value)
    if not txt:
        return None
    found = by_name.get(txt.lower())
    if found:
        return found
    return by_code.get(normalize_code(txt))


def _get_file_or_redirect():
    file = request.files.get("file")
    if not file or file.filename == "":
        flash("Veuillez sélectionner un fichier Excel (.xlsx) ou CSV.", "danger")
        return None
    return file


def _result(kind, errors=None, created=0, updated=0, warnings=None, details=None):
    return render_template(
        "excel_import_result.html", kind=kind, errors=errors or [], created=created,
        updated=updated, warnings=warnings or [], details=details or [],
    )


def _style_sheet(ws, headers, widths, number_formats=None):
    ws.append(headers)
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 20
    ws._rm_number_formats = number_formats or {}


def _apply_formats(ws):
    fmts = getattr(ws, "_rm_number_formats", {})
    for col_letter, fmt in fmts.items():
        for cell in ws[col_letter][1:]:
            cell.number_format = fmt


def _examples(ws, rows):
    """Lignes d'exemple (en gris) : ignorées à l'import car elles commencent
    par « Exemple » — pas de risque d'importer de fausses données."""
    for r in rows:
        ws.append(r)
        for cell in ws[ws.max_row]:
            cell.font = EXAMPLE_FONT


def _add_instructions(wb, lines):
    ws = wb.create_sheet("Mode d'emploi")
    ws.column_dimensions["A"].width = 110
    ws.append(["Mode d'emploi"])
    ws["A1"].font = Font(bold=True, size=14, color="6E3614")
    for line in lines:
        ws.append([line])
        ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True)


def _send_workbook(wb, filename):
    for ws in wb.worksheets:
        _apply_formats(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf, as_attachment=True, download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _stamp():
    return now_local().strftime("%Y%m%d_%H%M")


def _local_date(dt):
    """Pour l'export : date locale sans heure (affichée comme date dans Excel)."""
    return to_local(dt).replace(tzinfo=None) if dt else None


def _import_date_to_utc(raw):
    """Date lue dans un fichier = date LOCALE → stockée en UTC (à midi, pour
    que le jour affiché reste le bon)."""
    d = parse_date_any(raw)
    if d is None:
        return None
    if d.hour == 0 and d.minute == 0:
        d = d.replace(hour=12)
    return local_to_utc(d)


# ===========================================================================
# Page centrale
# ===========================================================================
@excel_bp.route("/")
@login_required
@roles_required(ROLE_ADMIN)
def excel_home():
    return render_template("excel_home.html")


# ===========================================================================
# INGRÉDIENTS / STOCK
# ===========================================================================
ING_HEADERS = ["nom_ingredient", "code", "catégorie", "unité", "stock", "seuil_alerte", "prix_unitaire"]
ING_WIDTHS = [24, 9, 16, 10, 12, 14, 15]
ING_FORMATS = {"E": "#,##0.###", "F": "#,##0.###", "G": "#,##0"}


@excel_bp.route("/template/ingredients")
@login_required
@roles_required(ROLE_ADMIN)
def template_ingredients():
    wb = Workbook()
    ws = wb.active
    ws.title = "Ingrédients"
    _style_sheet(ws, ING_HEADERS, ING_WIDTHS, ING_FORMATS)
    _examples(ws, [
        ["Exemple : Riz", "RIZ", "Céréale", "kg", 50, 10, 2500],
        ["Exemple : Poulet", "PLT", "Viande", "kg", 20, 5, 15000],
        ["Exemple : Yaourt", "", "Laitier", "pièce", 30, 6, 1200],
    ])
    _add_instructions(wb, [
        "• Une ligne = un ingrédient. Seul « nom_ingredient » est obligatoire.",
        "• « code » : code court (ex : PLT pour Poulet). Laissez vide → il est créé automatiquement.",
        "• « stock » : quantité ACTUELLE. Pour un ingrédient déjà existant, le stock est REMPLACÉ par cette valeur "
        "(et c'est noté dans le Journal du stock). Laissez la case vide pour ne pas toucher au stock.",
        "• Pour AJOUTER du stock acheté, utilisez plutôt l'import des Achats.",
        "• Nombres : 2500, 2 500, 0,3 ou 0.3 sont acceptés.",
        "• Les lignes qui commencent par « Exemple » sont ignorées : écrivez vos vraies lignes en dessous "
        "(ou effacez les exemples).",
    ])
    return _send_workbook(wb, "modele_ingredients.xlsx")


@excel_bp.route("/export/ingredients")
@login_required
@roles_required(ROLE_ADMIN)
def export_ingredients():
    wb = Workbook()
    ws = wb.active
    ws.title = "Ingrédients"
    _style_sheet(ws, ING_HEADERS, ING_WIDTHS, ING_FORMATS)
    for ing in Ingredient.query.order_by(Ingredient.name).all():
        ws.append([ing.name, ing.code, ing.category, ing.unit, ing.stock_quantity, ing.alert_threshold, ing.unit_price])
    return _send_workbook(wb, f"export_ingredients_{_stamp()}.xlsx")


@excel_bp.route("/import/ingredients", methods=["POST"])
@login_required
@roles_required(ROLE_ADMIN)
def import_ingredients():
    file = _get_file_or_redirect()
    if not file:
        return redirect(url_for("excel.excel_home"))
    try:
        mapping, rows = _read_table(file, "ingredients")
    except ImportFileError as e:
        return _result("ingrédients", errors=[str(e)])

    errors, warnings, parsed = [], [], {}
    by_name, by_code = _ingredient_lookup()
    codes_in_file = {}

    for n, row in rows:
        name = clean_text(_get(row, mapping, "name"))
        if not name:
            errors.append(f"Ligne {n} : nom de l'ingrédient manquant.")
            continue
        values = {}
        for key, label in (("stock", "stock"), ("threshold", "seuil_alerte"), ("price", "prix_unitaire")):
            raw = _get(row, mapping, key)
            if clean_text(raw) == "":
                values[key] = None
                continue
            num = parse_float(raw)
            if num is None:
                errors.append(f"Ligne {n} ({name}) : « {label} » doit être un nombre (valeur lue : {clean_text(raw)}).")
                break
            if num < 0:
                errors.append(f"Ligne {n} ({name}) : « {label} » ne peut pas être négatif.")
                break
            values[key] = num
        else:
            code = normalize_code(_get(row, mapping, "code"))
            if code:
                other = codes_in_file.get(code)
                if other and other.lower() != name.lower():
                    errors.append(f"Ligne {n} : le code « {code} » est déjà donné à « {other} » plus haut dans le fichier.")
                    continue
                owner = by_code.get(code)
                if owner and owner.name.strip().lower() != name.lower():
                    errors.append(f"Ligne {n} : le code « {code} » appartient déjà à l'ingrédient « {owner.name} ».")
                    continue
                codes_in_file[code] = name
            key = name.lower()
            if key in parsed:
                warnings.append(f"Ligne {n} : « {name} » apparaît plusieurs fois — la dernière ligne est gardée.")
            parsed[key] = {
                "name": name, "code": code,
                "category": clean_text(_get(row, mapping, "category")),
                "unit": clean_text(_get(row, mapping, "unit")),
                **values,
            }

    if errors:
        return _result("ingrédients", errors=errors)
    if not parsed:
        return _result("ingrédients", errors=["Aucune ligne à importer dans le fichier."])

    taken = set(by_code.keys()) | set(codes_in_file.keys())
    created = updated = 0
    details = []
    for key, d in parsed.items():
        ing = by_name.get(key) or (by_code.get(d["code"]) if d["code"] else None)
        if ing:
            if d["code"] and d["code"] != (ing.code or ""):
                ing.code = d["code"]
            if d["category"]:
                ing.category = d["category"]
            if d["unit"]:
                ing.unit = d["unit"]
            if d["threshold"] is not None:
                ing.alert_threshold = d["threshold"]
            if d["price"] is not None:
                ing.unit_price = d["price"]
            if d["stock"] is not None and abs(d["stock"] - ing.stock_quantity) > 1e-9:
                db.session.add(StockAdjustment(
                    ingredient_id=ing.id, old_quantity=ing.stock_quantity, new_quantity=d["stock"],
                    reason="Import Excel", changed_by_id=current_user.id,
                ))
                details.append(f"Stock « {ing.name} » : {ing.stock_quantity:g} → {d['stock']:g} {ing.unit}")
                ing.stock_quantity = d["stock"]
            updated += 1
        else:
            code = d["code"] or unique_code(suggest_code(d["name"]), taken)
            taken.add(code)
            db.session.add(Ingredient(
                name=d["name"], code=code, category=d["category"], unit=d["unit"] or "pièce",
                stock_quantity=d["stock"] or 0, alert_threshold=d["threshold"] or 0,
                unit_price=d["price"] or 0,
            ))
            details.append(f"Nouvel ingrédient : {d['name']} ({code})")
            created += 1

    db.session.commit()
    return _result("ingrédients", created=created, updated=updated, warnings=warnings, details=details)


# ===========================================================================
# MENU + RECETTE
# ===========================================================================
MENU_HEADERS = ["nom_menu", "catégorie_menu", "prix_vente", "ingredient", "quantité_par_portion"]
MENU_WIDTHS = [24, 16, 13, 22, 22]
MENU_FORMATS = {"C": "#,##0", "E": "#,##0.###"}


@excel_bp.route("/template/menu")
@login_required
@roles_required(ROLE_ADMIN)
def template_menu():
    wb = Workbook()
    ws = wb.active
    ws.title = "Menu"
    _style_sheet(ws, MENU_HEADERS, MENU_WIDTHS, MENU_FORMATS)
    _examples(ws, [
        ["Exemple : Poulet Riz", "Plat", 12000, "Riz", 0.3],
        ["Exemple : Poulet Riz", "Plat", 12000, "PLT", 0.25],
        ["Exemple : Jus naturel", "Boisson", 3000, "", ""],
    ])
    _add_instructions(wb, [
        "• Une ligne par ingrédient de la recette. Répétez le nom du menu sur chaque ligne "
        "(ou laissez « nom_menu » vide : la ligne est alors rattachée au menu du dessus).",
        "• « ingredient » : le NOM (ex : Poulet) ou le CODE (ex : PLT). L'ingrédient doit exister dans le Stock.",
        "• Un menu sans recette (ex : boisson achetée toute prête) : une seule ligne avec ingredient vide.",
        "• Si le menu existe déjà : prix et catégorie sont mis à jour ; la recette est remplacée "
        "seulement si le fichier contient des ingrédients pour ce menu.",
        "• Les photos des menus s'ajoutent dans le logiciel (page Menus), pas dans Excel.",
        "• Les lignes qui commencent par « Exemple » sont ignorées.",
    ])
    return _send_workbook(wb, "modele_menu.xlsx")


@excel_bp.route("/export/menu")
@login_required
@roles_required(ROLE_ADMIN)
def export_menu():
    wb = Workbook()
    ws = wb.active
    ws.title = "Menu"
    _style_sheet(ws, MENU_HEADERS, MENU_WIDTHS, MENU_FORMATS)
    for menu in MenuItem.active().order_by(MenuItem.name).all():
        if menu.recipe_items:
            for ri in menu.recipe_items:
                ws.append([menu.name, menu.category, menu.sale_price, ri.ingredient.name, ri.quantity_per_portion])
        else:
            ws.append([menu.name, menu.category, menu.sale_price, "", ""])
    return _send_workbook(wb, f"export_menu_{_stamp()}.xlsx")


@excel_bp.route("/import/menu", methods=["POST"])
@login_required
@roles_required(ROLE_ADMIN)
def import_menu():
    file = _get_file_or_redirect()
    if not file:
        return redirect(url_for("excel.excel_home"))
    try:
        mapping, rows = _read_table(file, "menu")
    except ImportFileError as e:
        return _result("menu", errors=[str(e)])

    errors, warnings = [], []
    groups = {}  # clé minuscule → {name, category, price, lines{ing_id: qty}}
    by_name, by_code = _ingredient_lookup()
    current = None

    for n, row in rows:
        name = clean_text(_get(row, mapping, "name"))
        ing_raw = _get(row, mapping, "ingredient")
        if not name:
            if current and clean_text(ing_raw):
                name = current["name"]  # ligne rattachée au menu du dessus
            else:
                errors.append(f"Ligne {n} : nom du menu manquant.")
                continue

        group = groups.get(name.lower())
        if group is None:
            group = groups[name.lower()] = {"name": name, "category": "", "price": None, "lines": {}}
        current = group

        cat = clean_text(_get(row, mapping, "category"))
        if cat:
            group["category"] = cat
        price_raw = _get(row, mapping, "price")
        if clean_text(price_raw) != "":
            price = parse_float(price_raw)
            if price is None or price < 0:
                errors.append(f"Ligne {n} ({name}) : prix_vente invalide (valeur lue : {clean_text(price_raw)}).")
                continue
            if group["price"] is not None and abs(group["price"] - price) > 1e-9:
                warnings.append(f"Ligne {n} : « {name} » a deux prix différents — {price:g} Ar est gardé.")
            group["price"] = price

        if clean_text(ing_raw):
            ing = _find_ingredient(ing_raw, by_name, by_code)
            if not ing:
                errors.append(
                    f"Ligne {n} ({name}) : ingrédient « {clean_text(ing_raw)} » introuvable dans le Stock "
                    "(vérifiez l'orthographe, utilisez le code, ou créez-le d'abord)."
                )
                continue
            qty = parse_float(_get(row, mapping, "qty"))
            if qty is None or qty <= 0:
                errors.append(f"Ligne {n} ({name}) : quantité_par_portion manquante ou invalide pour « {ing.name} ».")
                continue
            if ing.id in group["lines"]:
                warnings.append(f"Ligne {n} : « {ing.name} » est deux fois dans « {name} » — quantités additionnées.")
                group["lines"][ing.id] += qty
            else:
                group["lines"][ing.id] = qty

    if errors:
        return _result("menu", errors=errors)
    if not groups:
        return _result("menu", errors=["Aucune ligne à importer dans le fichier."])

    existing = {m.name.strip().lower(): m for m in MenuItem.query.all()}
    created = updated = 0
    details = []
    for key, g in groups.items():
        menu = existing.get(key)
        if menu:
            if g["category"]:
                menu.category = g["category"]
            if g["price"] is not None:
                menu.sale_price = g["price"]
            if g["lines"]:
                RecipeItem.query.filter_by(menu_item_id=menu.id).delete()
            if menu.is_archived:
                menu.is_archived = False
                details.append(f"Menu restauré : {menu.name}")
            updated += 1
        else:
            if g["price"] is None:
                warnings.append(f"« {g['name']} » créé sans prix (0 Ar) — pensez à le compléter.")
            menu = MenuItem(name=g["name"], category=g["category"], sale_price=g["price"] or 0)
            db.session.add(menu)
            db.session.flush()
            created += 1
            details.append(f"Nouveau menu : {g['name']}")
        for ing_id, qty in g["lines"].items():
            db.session.add(RecipeItem(menu_item_id=menu.id, ingredient_id=ing_id, quantity_per_portion=qty))

    db.session.commit()
    return _result("menus", created=created, updated=updated, warnings=warnings, details=details)


# ===========================================================================
# ACHATS
# ===========================================================================
ACHAT_HEADERS = ["date", "fournisseur", "ingredient", "quantité", "prix_unitaire", "source_financement"]
ACHAT_WIDTHS = [13, 22, 22, 12, 14, 22]
ACHAT_FORMATS = {"A": "DD/MM/YYYY", "D": "#,##0.###", "E": "#,##0"}


@excel_bp.route("/template/achats")
@login_required
@roles_required(ROLE_ADMIN)
def template_achats():
    wb = Workbook()
    ws = wb.active
    ws.title = "Achats"
    _style_sheet(ws, ACHAT_HEADERS, ACHAT_WIDTHS, ACHAT_FORMATS)
    _examples(ws, [
        ["Exemple", "Marché Analakely", "Riz", 20, 2400, ""],
        ["Exemple", "Boucherie", "PLT", 5, 15000, ""],
    ])
    sources = [s.name for s in FundingSource.query.order_by(FundingSource.name).all()]
    _add_instructions(wb, [
        "• Une ligne = un achat. Obligatoires : ingredient et quantité.",
        "• « ingredient » : le NOM ou le CODE (ex : PLT). Il doit déjà exister dans le Stock.",
        "• Le stock de l'ingrédient AUGMENTE de la quantité achetée.",
        "• « date » : vide = aujourd'hui. Formats acceptés : 23/09/2026, 2026-09-23, ou une date Excel.",
        "• « source_financement » (optionnel) doit exister dans Finances. Sources actuelles : "
        + (", ".join(sources) if sources else "aucune pour le moment") + ".",
        "• Attention : n'importez pas deux fois le même fichier d'achats (les achats seraient comptés deux fois).",
        "• Les lignes qui commencent par « Exemple » sont ignorées.",
    ])
    return _send_workbook(wb, "modele_achats.xlsx")


@excel_bp.route("/export/achats")
@login_required
@roles_required(ROLE_ADMIN)
def export_achats():
    wb = Workbook()
    ws = wb.active
    ws.title = "Achats"
    _style_sheet(ws, ACHAT_HEADERS + ["total", "statut", "saisi_par"], ACHAT_WIDTHS + [14, 12, 14],
                 {**ACHAT_FORMATS, "G": "#,##0"})
    for p in Purchase.query.order_by(Purchase.date.desc()).all():
        ws.append([
            _local_date(p.date), p.supplier, p.ingredient.name, p.quantity, p.unit_price,
            p.funding_source.name if p.funding_source else "", p.total,
            "Annulé" if p.is_cancelled else "Valide",
            p.created_by.username if p.created_by else "",
        ])
    return _send_workbook(wb, f"export_achats_{_stamp()}.xlsx")


@excel_bp.route("/import/achats", methods=["POST"])
@login_required
@roles_required(ROLE_ADMIN)
def import_achats():
    file = _get_file_or_redirect()
    if not file:
        return redirect(url_for("excel.excel_home"))
    try:
        mapping, rows = _read_table(file, "achats")
    except ImportFileError as e:
        return _result("achats", errors=[str(e)])

    errors, parsed = [], []
    by_name, by_code = _ingredient_lookup()
    source_names = {s.name.strip().lower(): s for s in FundingSource.query.all()}

    for n, row in rows:
        ing_raw = _get(row, mapping, "ingredient")
        if not clean_text(ing_raw):
            errors.append(f"Ligne {n} : ingredient manquant.")
            continue
        ing = _find_ingredient(ing_raw, by_name, by_code)
        if not ing:
            errors.append(f"Ligne {n} : ingrédient « {clean_text(ing_raw)} » introuvable dans le Stock "
                          "(vérifiez l'orthographe ou utilisez le code).")
            continue
        qty = parse_float(_get(row, mapping, "qty"))
        if qty is None or qty <= 0:
            errors.append(f"Ligne {n} ({ing.name}) : quantité doit être un nombre supérieur à 0.")
            continue
        price_raw = _get(row, mapping, "price")
        price = parse_float(price_raw, default=0) if clean_text(price_raw) else 0
        if price is None or price < 0:
            errors.append(f"Ligne {n} ({ing.name}) : prix_unitaire invalide.")
            continue

        date_raw = _get(row, mapping, "date")
        if clean_text(date_raw):
            date_val = _import_date_to_utc(date_raw)
            if date_val is None:
                errors.append(f"Ligne {n} : date « {clean_text(date_raw)} » non reconnue (utilisez JJ/MM/AAAA).")
                continue
        else:
            date_val = datetime.utcnow()

        source_id = None
        src_raw = clean_text(_get(row, mapping, "source"))
        if src_raw:
            src = source_names.get(src_raw.lower())
            if not src:
                errors.append(f"Ligne {n} : source de financement « {src_raw} » introuvable (voir page Finances).")
                continue
            source_id = src.id

        parsed.append({
            "date": date_val, "supplier": clean_text(_get(row, mapping, "supplier")),
            "ingredient": ing, "quantity": qty, "unit_price": price, "funding_source_id": source_id,
        })

    if errors:
        return _result("achats", errors=errors)
    if not parsed:
        return _result("achats", errors=["Aucune ligne à importer dans le fichier."])

    details = []
    for d in parsed:
        ing = d.pop("ingredient")
        db.session.add(Purchase(ingredient_id=ing.id, created_by_id=current_user.id, **d))
        ing.stock_quantity += d["quantity"]
        details.append(f"{ing.name} : +{d['quantity']:g} {ing.unit}")
    db.session.commit()
    return _result("achats", created=len(parsed), details=details)


# ===========================================================================
# DÉPENSES (fixes / variables)
# ===========================================================================
DEP_HEADERS = ["date", "type", "catégorie", "montant", "description", "source_financement"]
DEP_WIDTHS = [13, 11, 22, 14, 34, 22]
DEP_FORMATS = {"A": "DD/MM/YYYY", "D": "#,##0"}


@excel_bp.route("/template/depenses")
@login_required
@roles_required(ROLE_ADMIN)
def template_depenses():
    wb = Workbook()
    ws = wb.active
    ws.title = "Dépenses"
    _style_sheet(ws, DEP_HEADERS, DEP_WIDTHS, DEP_FORMATS)
    _examples(ws, [
        ["Exemple", "variable", "Huile (menaka)", 18000, "2 bidons", ""],
        ["Exemple", "fixe", "Loyer", 500000, "Loyer du mois", ""],
    ])
    _add_instructions(wb, [
        "• Une ligne = une dépense (hors achat d'ingrédient). Obligatoire : montant.",
        "• « type » : fixe ou variable (vide = variable).",
        "• « date » : vide = aujourd'hui.",
        "• N'importez pas deux fois le même fichier (les dépenses seraient comptées deux fois).",
        "• Les lignes qui commencent par « Exemple » sont ignorées.",
    ])
    return _send_workbook(wb, "modele_depenses.xlsx")


@excel_bp.route("/export/depenses")
@login_required
@roles_required(ROLE_ADMIN)
def export_depenses():
    wb = Workbook()
    ws = wb.active
    ws.title = "Dépenses"
    _style_sheet(ws, DEP_HEADERS + ["statut", "saisi_par"], DEP_WIDTHS + [12, 14], DEP_FORMATS)
    for e in Expense.query.order_by(Expense.date.desc()).all():
        ws.append([
            _local_date(e.date), "fixe" if e.is_fixed else "variable", e.category, e.amount,
            e.description or (e.fixed_expense.name if e.fixed_expense else ""),
            e.funding_source.name if e.funding_source else "",
            "Annulé" if e.is_cancelled else "Valide",
            e.created_by.username if e.created_by else "",
        ])
    return _send_workbook(wb, f"export_depenses_{_stamp()}.xlsx")


@excel_bp.route("/import/depenses", methods=["POST"])
@login_required
@roles_required(ROLE_ADMIN)
def import_depenses():
    file = _get_file_or_redirect()
    if not file:
        return redirect(url_for("excel.excel_home"))
    try:
        mapping, rows = _read_table(file, "depenses")
    except ImportFileError as e:
        return _result("dépenses", errors=[str(e)])

    errors, parsed = [], []
    source_names = {s.name.strip().lower(): s for s in FundingSource.query.all()}
    for n, row in rows:
        amount = parse_float(_get(row, mapping, "amount"))
        if amount is None or amount <= 0:
            errors.append(f"Ligne {n} : montant manquant ou invalide.")
            continue
        t = _norm(_get(row, mapping, "type"))
        if t in ("", "variable", "var", "v"):
            etype = EXPENSE_TYPE_VARIABLE
        elif t in ("fixe", "fixed", "fix", "f"):
            etype = EXPENSE_TYPE_FIXED
        else:
            errors.append(f"Ligne {n} : type « {clean_text(_get(row, mapping, 'type'))} » inconnu (fixe ou variable).")
            continue
        date_raw = _get(row, mapping, "date")
        if clean_text(date_raw):
            d = _import_date_to_utc(date_raw)
            if d is None:
                errors.append(f"Ligne {n} : date « {clean_text(date_raw)} » non reconnue (utilisez JJ/MM/AAAA).")
                continue
        else:
            d = datetime.utcnow()
        source_id = None
        src_raw = clean_text(_get(row, mapping, "source"))
        if src_raw:
            src = source_names.get(src_raw.lower())
            if not src:
                errors.append(f"Ligne {n} : source de financement « {src_raw} » introuvable.")
                continue
            source_id = src.id
        parsed.append(Expense(
            date=d, expense_type=etype,
            period=to_local(d).strftime("%Y-%m") if etype == EXPENSE_TYPE_FIXED else None,
            category=clean_text(_get(row, mapping, "category")) or "Autre",
            description=clean_text(_get(row, mapping, "description")) or None,
            amount=amount, funding_source_id=source_id, created_by_id=current_user.id,
        ))

    if errors:
        return _result("dépenses", errors=errors)
    if not parsed:
        return _result("dépenses", errors=["Aucune ligne à importer dans le fichier."])
    db.session.add_all(parsed)
    db.session.commit()
    return _result("dépenses", created=len(parsed))


# ===========================================================================
# VENTES (export seulement — les ventes se font à la caisse)
# ===========================================================================
@excel_bp.route("/export/ventes")
@login_required
@roles_required(ROLE_ADMIN)
def export_ventes():
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventes (détail)"
    _style_sheet(
        ws,
        ["n°_ticket", "date", "heure", "table", "servi_par", "client", "paiement", "menu", "quantité", "prix_unitaire", "total_ligne"],
        [12, 12, 8, 12, 14, 18, 15, 24, 10, 13, 13],
        {"B": "DD/MM/YYYY", "J": "#,##0", "K": "#,##0"},
    )
    ws2 = wb.create_sheet("Ventes par jour")
    _style_sheet(ws2, ["date", "nombre_tickets", "chiffre_affaires"], [13, 16, 18], {"A": "DD/MM/YYYY", "C": "#,##0"})

    per_day = {}
    for s in Sale.query.order_by(Sale.created_at.desc()).all():
        local = to_local(s.created_at)
        for item in s.items:
            ws.append([
                s.receipt_number(), local.replace(hour=0, minute=0, second=0, microsecond=0),
                local.strftime("%H:%M"), s.table_number or "", s.server.username if s.server else "",
                s.client_name or "", s.payment_method, item.menu_item.name if item.menu_item else "?",
                item.quantity, item.unit_price, item.line_total,
            ])
        day = local.date()
        cnt, tot = per_day.get(day, (0, 0))
        per_day[day] = (cnt + 1, tot + s.total)
    for day in sorted(per_day, reverse=True):
        cnt, tot = per_day[day]
        ws2.append([datetime(day.year, day.month, day.day), cnt, tot])
    return _send_workbook(wb, f"export_ventes_{_stamp()}.xlsx")
